"""Inventory inspection, loading, normalization, and validation."""

from __future__ import annotations

import csv
import json
import unicodedata
from zipfile import BadZipFile
from collections.abc import Iterable, Mapping
from datetime import date, datetime
from pathlib import Path
from typing import Any

import yaml
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from .config import CANONICAL_FIELDS, OPTIONAL_FIELDS, InventoryConfig
from .identity import normalize_asset
from .errors import InventoryValidationError, ValidationIssue
from .models import Asset

SUPPORTED_FORMATS = {"xlsx", "csv", "json", "yaml"}


def detect_format(path: Path, configured: str = "auto") -> str:
    if configured != "auto":
        value = configured.casefold().lstrip(".")
        if value == "yml":
            value = "yaml"
        if value not in SUPPORTED_FORMATS:
            raise InventoryValidationError(
                [ValidationIssue(f"unsupported inventory format: {configured}")]
            )
        return value
    suffix = path.suffix.casefold().lstrip(".")
    if suffix == "yml":
        suffix = "yaml"
    if suffix not in SUPPORTED_FORMATS:
        raise InventoryValidationError(
            [ValidationIssue(f"cannot detect inventory format from extension: {path.suffix}")]
        )
    return suffix


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return str(value)


def _normalized(value: Any) -> str:
    text = unicodedata.normalize("NFKC", _cell_text(value)).strip()
    if "\x00" in text:
        raise ValueError("contains a null character")
    return text


def _resolve_records(root: Any, path: str | None) -> Any:
    current = root
    if path:
        for component in path.split("."):
            if not component:
                raise KeyError("records path contains an empty component")
            if not isinstance(current, Mapping) or component not in current:
                raise KeyError(f"records path component not found: {component}")
            current = current[component]
    return current


def _validate_rows(rows: Iterable[tuple[str, Mapping[str, Any]]], columns: Mapping[str, str]) -> list[Asset]:
    issues: list[ValidationIssue] = []
    assets: list[Asset] = []
    seen_ids: dict[str, str] = {}
    for location, row in rows:
        if not isinstance(row, Mapping):
            issues.append(ValidationIssue("record must be an object", location))
            continue
        values: dict[str, str] = {}
        mapped = {name: columns.get(name, name) for name in CANONICAL_FIELDS + OPTIONAL_FIELDS}
        advanced = any(row.get(mapped[name]) for name in ("purl", "cpe", "ecosystem", "commit"))
        required = ("asset_id",) if advanced else CANONICAL_FIELDS
        if any(mapped[name] not in row for name in required):
            issues.append(ValidationIssue("record is missing mapped fields", location))
            continue
        row_has_value = False
        for canonical in CANONICAL_FIELDS + OPTIONAL_FIELDS:
            source_name = mapped[canonical]
            raw = row.get(source_name)
            if raw not in (None, ""):
                row_has_value = True
            try:
                if raw is not None and not isinstance(raw, str):
                    raise ValueError("must be text; numeric, date, boolean and compound values cannot preserve inventory spelling")
                # Exact package identifiers must not undergo compatibility
                # Unicode normalization (which can change registry identity).
                values[canonical] = _normalized(raw) if canonical in {"asset_id", "category", "system_id"} or (not advanced and canonical in {"vendor", "product"}) else (raw or "").strip()
                if "\x00" in values[canonical]:
                    raise ValueError("contains a null character")
            except (TypeError, ValueError) as exc:
                issues.append(
                    ValidationIssue(f"{canonical} cannot be normalized: {exc}", location)
                )
                values[canonical] = ""
        if not row_has_value:
            continue
        try:
            asset = normalize_asset(Asset(**values))
        except ValueError as exc:
            issues.append(ValidationIssue(str(exc), location))
            continue
        duplicate_key = values["asset_id"].casefold()
        if duplicate_key in seen_ids:
            issues.append(
                ValidationIssue(
                    f"duplicate asset_id; first seen at {seen_ids[duplicate_key]}", location
                )
            )
            continue
        seen_ids[duplicate_key] = location
        assets.append(asset)
    if issues:
        raise InventoryValidationError(issues)
    if not assets:
        raise InventoryValidationError([ValidationIssue("inventory contains no valid records")])
    return assets


def _xlsx_rows(config: InventoryConfig) -> Iterable[tuple[str, Mapping[str, Any]]]:
    try:
        workbook = load_workbook(config.path, read_only=True, data_only=False)
    except (OSError, ValueError, BadZipFile, InvalidFileException) as exc:
        raise InventoryValidationError(
            [ValidationIssue(f"cannot read XLSX workbook: {exc}")]
        ) from exc
    try:
        if config.worksheet:
            if config.worksheet not in workbook.sheetnames:
                raise InventoryValidationError(
                    [ValidationIssue(f"worksheet not found: {config.worksheet}")]
                )
            worksheet = workbook[config.worksheet]
        else:
            worksheet = workbook[workbook.sheetnames[0]]
        header_cells = next(
            worksheet.iter_rows(
                min_row=config.header_row, max_row=config.header_row, values_only=True
            ),
            None,
        )
        if header_cells is None:
            raise InventoryValidationError(
                [ValidationIssue(f"header row {config.header_row} does not exist")]
            )
        headers = [_normalized(value) for value in header_cells]
        if any(not value for value in headers):
            raise InventoryValidationError(
                [ValidationIssue("header row contains a blank column name")]
            )
        if len({value.casefold() for value in headers}) != len(headers):
            raise InventoryValidationError(
                [ValidationIssue("header row contains duplicate column names")]
            )
        header_lookup = {name.casefold(): name for name in headers}
        mapped = {name: config.columns.get(name, name) for name in CANONICAL_FIELDS + OPTIONAL_FIELDS}
        required = ("asset_id",) if any(mapped[name].casefold() in header_lookup for name in ("purl", "cpe", "ecosystem", "commit")) else CANONICAL_FIELDS
        missing = [mapped[name] for name in required if mapped[name].casefold() not in header_lookup]
        if missing:
            raise InventoryValidationError(
                [ValidationIssue("mapped columns not found: " + ", ".join(missing))]
            )
        for row_number, values in enumerate(
            worksheet.iter_rows(min_row=config.header_row + 1),
            start=config.header_row + 1,
        ):
            raw_cells = dict(zip(headers, values))
            if any(raw_cells[header_lookup[name.casefold()]].data_type == "f" for name in mapped.values() if name.casefold() in header_lookup):
                raise InventoryValidationError([ValidationIssue("mapped inventory cells must contain text, not formulas", f"row {row_number}")])
            raw = {name: cell.value for name, cell in raw_cells.items()}
            yield f"sheet {worksheet.title!r}, row {row_number}", {
                configured: raw.get(header_lookup[configured.casefold()])
                for configured in mapped.values() if configured.casefold() in header_lookup
            }
    finally:
        workbook.close()


def _csv_rows(config: InventoryConfig) -> Iterable[tuple[str, Mapping[str, Any]]]:
    try:
        handle = config.path.open("r", encoding=config.encoding, newline="")
    except (OSError, LookupError) as exc:
        raise InventoryValidationError([ValidationIssue(f"cannot read CSV: {exc}")]) from exc
    with handle:
        try:
            reader = csv.DictReader(handle, delimiter=config.delimiter, strict=True)
            if reader.fieldnames is None:
                raise InventoryValidationError([ValidationIssue("CSV has no header row")])
            headers = [_normalized(value) for value in reader.fieldnames]
            if len({value.casefold() for value in headers}) != len(headers):
                raise InventoryValidationError(
                    [ValidationIssue("CSV header contains duplicate column names")]
                )
            reader.fieldnames = headers
            header_lookup = {name.casefold(): name for name in headers}
            mapped = {name: config.columns.get(name, name) for name in CANONICAL_FIELDS + OPTIONAL_FIELDS}
            required = ("asset_id",) if any(mapped[name].casefold() in header_lookup for name in ("purl", "cpe", "ecosystem", "commit")) else CANONICAL_FIELDS
            missing = [mapped[name] for name in required if mapped[name].casefold() not in header_lookup]
            if missing:
                raise InventoryValidationError(
                    [ValidationIssue("mapped columns not found: " + ", ".join(missing))]
                )
            for row_number, row in enumerate(reader, start=2):
                if None in row:
                    raise InventoryValidationError([ValidationIssue("CSV row has more fields than its header", f"row {row_number}")])
                yield f"row {row_number}", {
                    configured: row.get(header_lookup[configured.casefold()])
                    for configured in mapped.values() if configured.casefold() in header_lookup
                }
        except csv.Error as exc:
            raise InventoryValidationError(
                [ValidationIssue(f"malformed CSV near line {reader.line_num}: {exc}")]
            ) from exc


def _structured_rows(config: InventoryConfig, kind: str) -> Iterable[tuple[str, Mapping[str, Any]]]:
    try:
        text = config.path.read_text(encoding=config.encoding)
        root = json.loads(text) if kind == "json" else yaml.safe_load(text)
        records = _resolve_records(root, config.records_path)
    except (OSError, LookupError, json.JSONDecodeError, yaml.YAMLError, KeyError) as exc:
        raise InventoryValidationError(
            [ValidationIssue(f"cannot read {kind.upper()} inventory: {exc}")]
        ) from exc
    if not isinstance(records, list):
        location = config.records_path or "document root"
        raise InventoryValidationError(
            [ValidationIssue("configured records value must be a list", location)]
        )
    for index, record in enumerate(records):
        yield f"record {index}", record


def load_inventory(config: InventoryConfig) -> list[Asset]:
    kind = detect_format(config.path, config.format)
    if kind == "xlsx":
        rows = _xlsx_rows(config)
    elif kind == "csv":
        rows = _csv_rows(config)
    else:
        rows = _structured_rows(config, kind)
    return _validate_rows(rows, config.columns)


def inspect_inventory(path: str | Path, *, header_row: int = 1, encoding: str = "utf-8-sig", delimiter: str = ",") -> dict[str, Any]:
    source = Path(path).expanduser().resolve()
    kind = detect_format(source)
    if kind == "xlsx":
        try:
            workbook = load_workbook(source, read_only=True, data_only=True)
        except (OSError, ValueError, BadZipFile, InvalidFileException) as exc:
            raise InventoryValidationError(
                [ValidationIssue(f"cannot read XLSX workbook: {exc}")]
            ) from exc
        try:
            sheets: list[dict[str, Any]] = []
            for worksheet in workbook.worksheets:
                headers = next(
                    worksheet.iter_rows(
                        min_row=header_row, max_row=header_row, values_only=True
                    ),
                    (),
                )
                sheets.append(
                    {
                        "name": worksheet.title,
                        "headers": [_cell_text(value).strip() for value in headers],
                        "max_row": worksheet.max_row,
                        "max_column": worksheet.max_column,
                    }
                )
            return {"path": str(source), "format": kind, "sheets": sheets}
        finally:
            workbook.close()
    if kind == "csv":
        try:
            with source.open("r", encoding=encoding, newline="") as handle:
                reader = csv.reader(handle, delimiter=delimiter)
                headers = next(reader, [])
                sample = [row for _, row in zip(range(3), reader)]
        except (OSError, LookupError, csv.Error) as exc:
            raise InventoryValidationError([ValidationIssue(f"cannot read CSV: {exc}")]) from exc
        return {"path": str(source), "format": kind, "headers": headers, "sample_rows": sample}
    try:
        text = source.read_text(encoding=encoding)
        root = json.loads(text) if kind == "json" else yaml.safe_load(text)
    except (OSError, LookupError, json.JSONDecodeError, yaml.YAMLError) as exc:
        raise InventoryValidationError(
            [ValidationIssue(f"cannot read {kind.upper()} inventory: {exc}")]
        ) from exc
    result: dict[str, Any] = {"path": str(source), "format": kind, "root_type": type(root).__name__}
    if isinstance(root, list):
        result["record_count"] = len(root)
        if root and isinstance(root[0], Mapping):
            result["first_record_fields"] = list(root[0].keys())
    elif isinstance(root, Mapping):
        result["top_level_keys"] = list(root.keys())
        result["list_keys"] = [key for key, value in root.items() if isinstance(value, list)]
    return result
