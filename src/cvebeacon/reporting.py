"""On-demand XLSX and JSON reporting."""

from __future__ import annotations

import json
import os
import re
import tempfile
from datetime import datetime, timezone
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import Applicability, QueryResult
from .errors import ReportingError


def _safe(value):
    if isinstance(value, str):
        value = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", lambda match: f"\\x{ord(match.group()):02x}", value)
        if len(value) > 32_000:
            value = value[:31_980] + "…[truncated]"
        if value.startswith(("=", "+", "-", "@")):
            return "'" + value
    return value


def _one_line(value):
    if isinstance(value, str):
        return " ".join(value.split())
    return value


def _detail_text(value: dict) -> str:
    text = json.dumps(value, sort_keys=True, ensure_ascii=False, separators=(",", ":"))
    return text if len(text) <= 32_000 else text[:31_980] + "…[truncated]"


def _excel_datetime(value):
    if isinstance(value, datetime) and value.tzinfo is not None:
        return value.astimezone(timezone.utc).replace(tzinfo=None)
    return value


def write_json(results: Iterable[QueryResult], path: str | Path) -> Path:
    destination = Path(path)
    temporary = None
    try:
        destination.parent.mkdir(parents=True, exist_ok=True)
        with tempfile.NamedTemporaryFile("w", encoding="utf-8", dir=destination.parent, suffix=".json", delete=False) as handle:
            temporary = Path(handle.name)
            json.dump([item.to_dict() for item in results], handle, indent=2, ensure_ascii=False)
        os.replace(temporary, destination)
    except (OSError, TypeError, ValueError) as exc:
        if temporary:
            temporary.unlink(missing_ok=True)
        raise ReportingError(f"cannot write JSON report {destination}: {exc}") from exc
    return destination


def _finish_sheet(sheet) -> None:
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.showGridLines = False
    used = sheet.iter_rows()
    for row in used:
        for cell in row:
            cell.font = Font(name="Arial", size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=False)
    for cell in sheet[1]:
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.row_dimensions[1].height = 30
    for column in range(1, sheet.max_column + 1):
        width = max((len(str(sheet.cell(row, column).value or "")) for row in range(1, min(sheet.max_row, 200) + 1)), default=10)
        sheet.column_dimensions[get_column_letter(column)].width = min(max(width + 2, 10), 55)


def write_xlsx(results: Iterable[QueryResult], path: str | Path) -> Path:
    values = list(results)
    destination = Path(path)
    try:
        destination.parent.mkdir(parents=True, exist_ok=True)
    except OSError as exc:
        raise ReportingError(f"cannot prepare report directory {destination.parent}: {exc}") from exc
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Summary"
    summary.append(["Asset ID", "Vendor", "Product", "Version", "Findings", "Affected", "Needs Review", "Coverage"])
    for result in values:
        summary.append([_safe(value) for value in [
            result.asset.asset_id, result.asset.vendor, result.asset.product, result.asset.version,
            len(result.findings), sum(x.applicability == Applicability.AFFECTED for x in result.findings),
            sum(x.applicability == Applicability.NEEDS_REVIEW for x in result.findings),
            result.coverage.value if result.coverage else "evaluated",
        ]])
    findings = workbook.create_sheet("Findings")
    findings.append(["Asset ID", "Vendor", "Product", "Version", "CVE", "Applicability", "Confidence", "Reason", "CVSS", "CVSS Vector", "EPSS", "EPSS Percentile", "EPSS Date", "CISA KEV", "EU KEV", "Rejected", "Published", "Modified", "Summary", "Sources", "Conflicts"])
    for result in values:
        for item in result.findings:
            vuln = item.vulnerability
            findings.append([_safe(_one_line(value)) for value in [
                item.asset.asset_id, item.asset.vendor, item.asset.product, item.asset.version,
                vuln.cve_id, item.applicability.value, item.confidence, item.reason,
                vuln.cvss_score, vuln.cvss_vector, vuln.epss_score, vuln.epss_percentile,
                vuln.epss_date.isoformat() if vuln.epss_date else None,
                vuln.cisa_kev, vuln.eu_kev, vuln.rejected, vuln.published, vuln.modified,
                vuln.summary, ", ".join(sorted({x.source for x in item.evidence})),
                "; ".join(item.conflicts),
            ]])
    uncertainty = workbook.create_sheet("Uncertainty")
    uncertainty.append(["Asset ID", "CVE", "State", "Reason", "Source Health"])
    for result in values:
        if result.coverage:
            uncertainty.append([_safe(_one_line(value)) for value in [result.asset.asset_id, "", result.coverage.value, result.coverage_reason, "; ".join(f"{x.source}:{x.status.value}" for x in result.source_health)]])
        for item in result.findings:
            if item.applicability in {Applicability.NEEDS_REVIEW, Applicability.COVERAGE_UNKNOWN}:
                uncertainty.append([_safe(_one_line(value)) for value in [item.asset.asset_id, item.vulnerability.cve_id, item.applicability.value, item.reason, ""]])
    evidence_sheet = workbook.create_sheet("Evidence")
    evidence_sheet.append(["Asset ID", "CVE", "Source", "Role", "Statement", "Source Timestamp", "Retrieved At", "Source URL", "Details"])
    for result in values:
        for item in result.findings:
            for evidence in item.evidence:
                evidence_sheet.append([_safe(_one_line(value)) for value in [
                    item.asset.asset_id, item.vulnerability.cve_id, evidence.source, evidence.role,
                    evidence.statement, evidence.source_timestamp, _excel_datetime(evidence.retrieved_at),
                    evidence.source_url, _detail_text(evidence.details),
                ]])
            for reference in item.vulnerability.references:
                evidence_sheet.append([_safe(_one_line(value)) for value in [
                    item.asset.asset_id, item.vulnerability.cve_id, "record", "reference",
                    "vulnerability reference", None, None, reference, "",
                ]])
    health_sheet = workbook.create_sheet("Source Health")
    health_sheet.append(["Asset ID", "Source", "Status", "Checked At", "Freshness At", "Message"])
    for result in values:
        for health in result.source_health:
            health_sheet.append([_safe(_one_line(value)) for value in [
                result.asset.asset_id, health.source, health.status.value, _excel_datetime(health.checked_at),
                _excel_datetime(health.freshness_at), health.message,
            ]])
    for sheet in workbook.worksheets:
        _finish_sheet(sheet)
    for row in findings.iter_rows(min_row=2, min_col=9, max_col=12):
        row[0].number_format = "0.0"
        row[2].number_format = "0.00000"
        row[3].number_format = "0.00000"
    for row in evidence_sheet.iter_rows(min_row=2, min_col=7, max_col=7):
        row[0].number_format = "yyyy-mm-dd hh:mm:ss"
    for row in health_sheet.iter_rows(min_row=2, min_col=4, max_col=5):
        row[0].number_format = "yyyy-mm-dd hh:mm:ss"
        row[1].number_format = "yyyy-mm-dd hh:mm:ss"
    temporary = None
    try:
        with tempfile.NamedTemporaryFile(dir=destination.parent, suffix=".xlsx", delete=False) as handle:
            temporary = Path(handle.name)
        workbook.save(temporary)
        os.replace(temporary, destination)
    except (OSError, TypeError, ValueError) as exc:
        if temporary:
            temporary.unlink(missing_ok=True)
        raise ReportingError(f"cannot write XLSX report {destination}: {exc}") from exc
    return destination
