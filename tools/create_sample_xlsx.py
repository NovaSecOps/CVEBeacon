"""Create the documented non-sensitive sample XLSX inventory."""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

destination = Path(__file__).resolve().parents[1] / "examples" / "inventory.xlsx"
workbook = Workbook()
sheet = workbook.active
sheet.title = "Assets"
sheet.append(["Asset ID", "Vendor", "Product", "Version"])
sheet.append(["asset-001", "Example Vendor", "Example Product", "1.0.0"])
sheet.freeze_panes = "A2"
sheet.auto_filter.ref = "A1:D2"
sheet.sheet_view.showGridLines = False
for cell in sheet[1]:
    cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="1F4E78")
    cell.alignment = Alignment(horizontal="center")
for width, column in zip((16, 20, 22, 14), ("A", "B", "C", "D")):
    sheet.column_dimensions[column].width = width
workbook.save(destination)
