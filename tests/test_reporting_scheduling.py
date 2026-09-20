from __future__ import annotations

from dataclasses import replace
from pathlib import Path

import pytest
from openpyxl import load_workbook

from cvebeacon.errors import SchedulingError
from cvebeacon.models import Applicability, Asset, Evidence, Finding, HealthStatus, QueryResult, SourceHealth, Vulnerability, utc_now
from cvebeacon.reporting import write_xlsx
from cvebeacon.scheduling import CRON_BEGIN, CRON_END, SchedulePlan, _windows_args, _without_managed_block, describe, install, make_plan, remove, status


def test_xlsx_has_practical_sheets(tmp_path):
    asset = Asset("a1", "Acme", "Widget", "1")
    finding = Finding(asset, Vulnerability("CVE-2026-1234"), Applicability.NEEDS_REVIEW, "limited", "unknown range", (Evidence("nvd", "test", "evidence"),))
    health = SourceHealth("nvd", HealthStatus.OK, utc_now(), "ok")
    path = write_xlsx([QueryResult(asset, (finding,), (health,))], tmp_path / "report.xlsx")
    workbook = load_workbook(path)
    assert workbook.sheetnames == ["Summary", "Findings", "Uncertainty", "Evidence", "Source Health", "Identities", "Advisory Details"]
    assert workbook["Findings"]["E2"].value == "CVE-2026-1234"


def test_xlsx_neutralizes_formula_like_inventory_values(tmp_path):
    asset = Asset("=HYPERLINK(\"bad\")", "+Vendor", "@Product", "-1")
    path = write_xlsx([QueryResult(asset, (), ())], tmp_path / "safe.xlsx")
    workbook = load_workbook(path, data_only=False)
    assert all(workbook["Summary"].cell(2, column).data_type != "f" for column in range(1, 5))


def test_advisory_report_preserves_identity_and_formula_safety(tmp_path):
    asset = Asset("a", product="requests", version="1", purl="pkg:pypi/requests@1", category="=1+1", system_id="@group")
    vuln = Vulnerability(advisory_id="GHSA-test-only", aliases=("PYSEC-2099-1",), fixed_versions=("+1",))
    finding = Finding(asset, vuln, Applicability.AFFECTED, "high", "range")
    workbook = load_workbook(write_xlsx([QueryResult(asset, (finding,), ())], tmp_path / "report.xlsx"))
    assert workbook["Findings"]["E2"].value == "GHSA-test-only"
    assert workbook["Identities"]["F2"].value == asset.purl
    assert workbook["Summary"]["I2"].value == "'=1+1"
    assert workbook["Advisory Details"]["D2"].value == "PYSEC-2099-1"
    assert not any(cell.data_type == "f" for sheet in workbook for row in sheet for cell in row)


@pytest.mark.parametrize("prefix", ["=", "+", "-", "@"])
def test_xlsx_neutralizes_upstream_strings_on_every_sheet(tmp_path, prefix):
    text = prefix + "SUM(1,2)"
    asset = Asset("a1", "Acme", "Widget", "1")
    evidence = Evidence("nvd", text, text, source_url=text, source_timestamp=text, details={"value": text})
    finding = Finding(asset, Vulnerability("CVE-2026-1234", summary=text, references=(text,)), Applicability.NEEDS_REVIEW, "limited", text, (evidence,), (text,))
    health = SourceHealth("nvd", HealthStatus.FAILED, utc_now(), text)
    path = write_xlsx([QueryResult(asset, (finding,), (health,), Applicability.COVERAGE_UNKNOWN, text)], tmp_path / "safe.xlsx")
    workbook = load_workbook(path, data_only=False)
    assert not any(cell.data_type == "f" for sheet in workbook for row in sheet for cell in row)
    assert len(workbook["Findings"]["A"]) == 2


def test_cron_removal_preserves_unrelated_entries():
    text = f"MAILTO=x\n{CRON_BEGIN}\n0 */4 * * * cvebeacon scan\n{CRON_END}\n15 2 * * * backup\n"
    clean = _without_managed_block(text)
    assert "MAILTO=x" in clean and "backup" in clean and "cvebeacon" not in clean


def test_duplicate_cron_blocks_are_rejected():
    text = f"{CRON_BEGIN}\na\n{CRON_END}\n{CRON_BEGIN}\nb\n{CRON_END}\n"
    with pytest.raises(SchedulingError): _without_managed_block(text)


def test_schedule_minimum_and_source_command(tmp_path, monkeypatch):
    with pytest.raises(SchedulingError): make_plan(tmp_path / "config.toml", 1)
    monkeypatch.setattr("cvebeacon.scheduling.platform.system", lambda: "Linux")
    plan = make_plan(tmp_path / "config.toml", 4, executable="/opt/cvebeacon")
    assert plan.command[-1] == "scan"
    assert "every 4 hours" in describe(plan)


def test_windows_24_hours_uses_daily_schedule(tmp_path, monkeypatch):
    monkeypatch.setattr("cvebeacon.scheduling.platform.system", lambda: "Windows")
    plan = make_plan(tmp_path / "config.toml", 24, executable="C:/Tools/cvebeacon.exe")
    args = _windows_args(plan)
    assert args[args.index("/SC") + 1] == "DAILY"


def test_incompatible_override_only_allowed_for_generation(tmp_path, monkeypatch):
    monkeypatch.setattr("cvebeacon.scheduling.platform.system", lambda: "Windows")
    with pytest.raises(SchedulingError): make_plan(tmp_path / "config.toml", 4, system="linux")
    assert make_plan(tmp_path / "config.toml", 4, system="linux", executable="/opt/cvebeacon", allow_incompatible=True).system == "linux"


def test_windows_schedule_commands_use_exact_owned_name(tmp_path, monkeypatch):
    monkeypatch.setattr("cvebeacon.scheduling.platform.system", lambda: "Windows")
    calls = []
    class Result:
        returncode = 0; stdout = "TaskName: CVEBeacon Monitor"; stderr = ""
    def fake_run(args, **kwargs): calls.append(args); return Result()
    plan = make_plan(tmp_path / "config.toml", 4, executable="C:/Tools/cvebeacon.exe")
    install(plan, run=fake_run)
    assert calls[0][calls[0].index("/TN") + 1] == "CVEBeacon Monitor"
    assert "CVEBeacon Monitor" in status(run=fake_run)
    assert remove(run=fake_run)
    assert all("*" not in value for call in calls for value in call)


def test_linux_install_preserves_unrelated_crontab(tmp_path, monkeypatch):
    monkeypatch.setattr("cvebeacon.scheduling.platform.system", lambda: "Linux")
    writes = []
    class Result:
        def __init__(self, returncode=0, stdout="", stderr=""):
            self.returncode, self.stdout, self.stderr = returncode, stdout, stderr
    def fake_run(args, **kwargs):
        if args == ["crontab", "-l"]: return Result(stdout="15 2 * * * backup\n")
        writes.append(kwargs["input"]); return Result()
    plan = make_plan(tmp_path / "config.toml", 4, executable="/opt/cvebeacon")
    install(plan, run=fake_run)
    assert "backup" in writes[0] and CRON_BEGIN in writes[0] and CRON_END in writes[0]
